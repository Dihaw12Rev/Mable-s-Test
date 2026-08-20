"""Render the analysis as an interactive HTML report and a Word document."""

from __future__ import annotations

import json
import re
from datetime import datetime
from pathlib import Path
from typing import Any

TEMPLATE = Path(__file__).parent / "templates" / "report.html"
DASHBOARD_SHELL = Path(__file__).parent / "templates" / "dashboard.html"
EXPLORER = Path(__file__).parent / "templates" / "explorer.html"
GUIDE = Path(__file__).parent / "templates" / "guide.html"
PLAN_SCREEN = Path(__file__).parent / "templates" / "plan.html"
PLAN_PRINT = Path(__file__).parent / "templates" / "plan-print.html"


# ------------------------------------------------------------------------- shared


def portal_name(analysis: dict[str, Any]) -> str:
    acct = analysis.get("account") or {}
    # uiDomain is "app.hubspot.com" for every portal — never a usable name.
    return str(acct.get("companyName") or f"Portal {analysis.get('portal_id', '')}")


def _date(value: Any) -> str:
    if not value:
        return "—"
    text = str(value)
    try:
        return datetime.fromisoformat(text.replace("Z", "+00:00")).date().isoformat()
    except ValueError:
        return text[:10]


def _strip_ticks(text: str) -> str:
    return re.sub(r"`([^`]+)`", r"\1", text or "")


# --------------------------------------------------------------------------- html


def render_html(analysis: dict[str, Any], out_path: Path) -> Path:
    template = TEMPLATE.read_text()
    payload = json.dumps(analysis, default=str)
    # Keep the JSON from terminating the host <script> element.
    payload = payload.replace("</", "<\\/")
    html = template.replace("__PORTAL_TITLE__", _escape(portal_name(analysis)))
    html = html.replace("__ANALYSIS_JSON__", payload)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(html)
    return out_path


def _escape(text: str) -> str:
    return (
        text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    )


# ----------------------------------------------------------------------- explorer


def render_explorer(analysis: dict[str, Any], out_path: Path) -> Path:
    """The interactive dependency explorer.

    Blast radius is computed in the browser rather than precomputed per node: the
    graph is small enough to traverse instantly, and it lets the reader change the
    hop depth without regenerating anything.
    """
    from . import graph as graph_mod

    g = analysis.get("graph") or graph_mod.build(analysis).to_dict()
    counts: dict[str, int] = {}
    for node in g["nodes"].values():
        counts[node["type"]] = counts.get(node["type"], 0) + 1

    order = ["workflow", "property", "list", "form", "email", "page", "pipeline", "stage"]
    stats = [(graph_mod.TYPE_PLURALS.get(t, t), counts[t]) for t in order if counts.get(t)]
    stats.append(("connections", len(g["edges"])))

    payload = {
        "portal_name": portal_name(analysis),
        "portal_id": analysis.get("portal_id"),
        "extracted_at": _date(analysis.get("extracted_at")),
        "graph": g,
        "stats": stats,
        "verbs": graph_mod.RELATIONS,
        "type_labels": graph_mod.TYPE_LABELS,
        "type_plurals": graph_mod.TYPE_PLURALS,
    }
    blob = json.dumps(payload, default=str).replace("</", "<\\/")
    html = (
        EXPLORER.read_text()
        .replace("__PORTAL_TITLE__", _escape(portal_name(analysis)))
        .replace("__PAYLOAD_JSON__", blob)
    )
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(html)
    return out_path


def render_guide(analysis: dict[str, Any], out_path: Path) -> Path:
    """The client-facing account guide.

    Three levels of disclosure — overview, business area, single asset — because a
    flat index of a few thousand assets is unusable to anyone who did not build the
    portal. The reader starts from a concept they recognise and drills in.
    """
    from . import areas as areas_mod
    from . import dashboard
    from . import graph as graph_mod

    g = graph_mod.build(analysis)
    counts: dict[str, int] = {}
    for node in g.nodes.values():
        counts[node["type"]] = counts.get(node["type"], 0) + 1
    order = ["workflow", "property", "list", "form", "email", "page", "pipeline", "stage"]

    payload = {
        "portal_name": portal_name(analysis),
        "portal_id": analysis.get("portal_id"),
        "extracted_at": _date(analysis.get("extracted_at")),
        "graph": g.to_dict(),
        "areas": areas_mod.build(analysis, g),
        "health": areas_mod.health(analysis),
        "findings": dashboard.findings(analysis),
        "stats": [(graph_mod.TYPE_PLURALS.get(t, t), counts[t]) for t in order if counts.get(t)],
        "verbs": graph_mod.RELATIONS,
        "type_labels": graph_mod.TYPE_LABELS,
        "type_plurals": graph_mod.TYPE_PLURALS,
    }
    blob = json.dumps(payload, default=str).replace("</", "<\\/")
    html = (
        GUIDE.read_text()
        .replace("__PORTAL_TITLE__", _escape(portal_name(analysis)))
        .replace("__PAYLOAD_JSON__", blob)
    )
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(html)
    return out_path


# ---------------------------------------------------------------------- review plan

RISK_WORD = {"safe": "safe", "judgment": "judgment", "careful": "highest risk",
             "process": "process"}


def _plan_body(analysis: dict[str, Any]) -> tuple[str, list[dict[str, Any]]]:
    """Render the phase blocks once; both the screen and print pages use them."""
    from . import graph as graph_mod
    from . import review

    g = graph_mod.build(analysis)
    work = review.cohorts(analysis, g)
    phases = review.phases(analysis, work)

    blocks = []
    for phase in phases:
        tags = []
        if phase["risk"] != "process":
            tags.append(f'<span class="tag {phase["risk"]}">{RISK_WORD[phase["risk"]]}</span>')
        if phase["count"]:
            tags.append(f'<span class="tag count">{_escape(phase["count"])}</span>')
        tags.append(f'<span class="tag count">{_escape(phase["effort"])}</span>')
        if phase["tab"]:
            tags.append(f'<span class="tag count">tab: {_escape(phase["tab"])}</span>')

        steps = "".join(f"<li>{_escape(t)}</li>" for t in phase["steps"])
        guard = (
            f'<div class="guard"><b>The mistake to avoid</b>{_escape(phase["guard"])}</div>'
            if phase["guard"] else ""
        )
        blocks.append(
            f'<section class="phase {phase["risk"]}" id="{phase["id"]}">'
            f'<div class="ph-head"><span class="ph-num">Phase {phase["n"]}</span>'
            f'<h2>{_escape(phase["title"])}</h2></div>'
            f'<div class="tagrow">{"".join(tags)}</div>'
            f'<p class="lead-q"><b>What this is.</b> {_escape(phase["what"])}</p>'
            f'<p class="why"><b>Why it matters.</b> {_escape(phase["why"])}</p>'
            f'<p class="why fixline"><b>How to fix it.</b> {_escape(phase["fix"])}</p>'
            f'<div class="where"><b>Where to work:</b> {_escape(phase["where"])}</div>'
            f'<div class="steplabel">Step by step</div>'
            f'<ol class="steps">{steps}</ol>{guard}</section>'
        )

    from . import review as review_mod

    glossary = (
        '<section class="phase glossary" id="glossary">'
        '<div class="ph-head"><span class="ph-num">Reference</span>'
        "<h2>What the words mean</h2></div>"
        '<p class="why">Every term this plan uses, in plain language.</p><dl>'
        + "".join(
            f"<dt>{_escape(term)}</dt><dd>{_escape(meaning)}</dd>"
            for term, meaning in review_mod.GLOSSARY
        )
        + "</dl></section>"
    )
    return "".join(blocks) + glossary, phases


def _plan_totals(analysis: dict[str, Any], phases) -> str:
    from . import review

    tiles = [("Workflows", f"{len(analysis['workflows']):,}")]
    for phase in phases:
        if phase["risk"] in (review.SAFE, review.CARE, review.JUDGE) and phase["tab"]:
            number = phase["count"].split()[0]
            tiles.append((phase["title"], number))
    return "".join(
        f'<div class="tot"><b>{_escape(v)}</b><span>{_escape(k)}</span></div>' for k, v in tiles
    )


def render_review_plan(analysis: dict[str, Any], out_html: Path, out_pdf: Path) -> list[Path]:
    """The review plan, as a screen page and a printable PDF from one set of phases."""
    body, phases = _plan_body(analysis)
    totals = _plan_totals(analysis, phases)
    nav = "".join(
        f'<li><a class="s-{p["risk"]}" href="#{p["id"]}">'
        f'<span class="n">{p["n"]}</span> {_escape(p["title"])}</a></li>'
        for p in phases
    )
    name = portal_name(analysis)
    written: list[Path] = []

    for template, target in ((PLAN_SCREEN, out_html), (PLAN_PRINT, out_pdf)):
        html = (
            template.read_text()
            .replace("__PORTAL__", _escape(name))
            .replace("__DATE__", _escape(_date(analysis.get("extracted_at"))))
            .replace("__TOTALS__", totals)
            .replace("__NAV__", nav)
            .replace("__BODY__", body)
        )
        target.parent.mkdir(parents=True, exist_ok=True)
        if target.suffix == ".pdf":
            from weasyprint import HTML

            HTML(string=html, base_url=str(PLAN_PRINT.parent)).write_pdf(str(target))
        else:
            target.write_text(html)
        written.append(target)
    return written


# ---------------------------------------------------------------------------- pdf


def render_dashboard_pdf(analysis: dict[str, Any], out_path: Path) -> Path:
    """The at-a-glance dashboard, rendered to PDF.

    WeasyPrint is used rather than a headless browser so the generator needs no
    browser install; every chart is server-rendered SVG for the same reason.
    """
    from weasyprint import HTML

    from . import dashboard

    name = portal_name(analysis)
    footer = f"{name} · HubSpot portal audit · {_date(analysis.get('extracted_at'))}"
    shell = DASHBOARD_SHELL.read_text()
    html = (
        shell.replace("__TITLE__", _escape(name))
        .replace("__FOOTER__", _escape(footer).replace('"', "'"))
        .replace("__BODY__", dashboard.build_html(analysis, name))
    )
    out_path.parent.mkdir(parents=True, exist_ok=True)
    HTML(string=html, base_url=str(DASHBOARD_SHELL.parent)).write_pdf(str(out_path))
    return out_path


# --------------------------------------------------------------------------- docx


def render_docx(analysis: dict[str, Any], out_path: Path) -> Path:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor, Inches

    doc = Document()
    _docx_styles(doc, Pt, RGBColor)

    name = portal_name(analysis)
    pid = analysis.get("portal_id", "")

    # -- cover -------------------------------------------------------------
    title = doc.add_paragraph(name, style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    doc.add_paragraph("HubSpot Portal Documentation", style="Subtitle")
    meta = doc.add_paragraph()
    meta.add_run(f"Portal ID {pid}   ·   Snapshot taken {_date(analysis.get('extracted_at'))}").italic = True
    doc.add_paragraph(
        "An as-built record of what this portal automates and which properties carry the "
        "load. Every asset below links back to itself in HubSpot."
    )

    # -- summary -----------------------------------------------------------
    doc.add_heading("At a glance", level=1)
    workflows = analysis["workflows"]
    props = analysis["properties"]
    rows = [
        ("Workflows", len(workflows)),
        ("  of which active", sum(1 for w in workflows if w["enabled"])),
        ("Properties", len(props)),
        ("  custom properties", sum(1 for p in props if not p["hubspot_defined"])),
        ("  unused custom properties", sum(1 for p in props if p["unused"])),
        ("Lists", len(analysis["lists"])),
        ("Forms", len(analysis["forms"])),
        ("Marketing emails", len(analysis["marketing_emails"])),
        ("Landing pages", len(analysis["landing_pages"])),
        ("Site pages", len(analysis["site_pages"])),
    ]
    table = doc.add_table(rows=0, cols=2)
    table.style = "Light Grid Accent 1"
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = str(value)

    # -- systems -----------------------------------------------------------
    doc.add_heading("Automation systems", level=1)
    doc.add_paragraph(
        "Workflows grouped by the job they do, inferred from their names and the properties "
        "they write. Start here to understand what the portal actually runs."
    )
    by_id = {w["id"]: w for w in workflows}
    for system, ids in analysis["systems"].items():
        members = [by_id[i] for i in ids if i in by_id]
        if not members:
            continue
        active = sum(1 for w in members if w["enabled"])
        doc.add_heading(system, level=2)
        written: dict[str, int] = {}
        for w in members:
            for p in w["properties_written"]:
                written[p] = written.get(p, 0) + 1
        top = sorted(written, key=lambda k: -written[k])[:8]
        summary = f"{len(members)} workflow(s), {active} active."
        if top:
            summary += " Writes to " + ", ".join(top) + "."
        doc.add_paragraph(summary)

        t = doc.add_table(rows=1, cols=4)
        t.style = "Light Grid Accent 1"
        for i, head in enumerate(["Workflow", "Status", "Writes", "Steps"]):
            t.rows[0].cells[i].text = head
        for w in sorted(members, key=lambda w: (not w["enabled"], w["name"].lower())):
            cells = t.add_row().cells
            _hyperlink(cells[0].paragraphs[0], w["link"], w["name"])
            cells[1].text = "Active" if w["enabled"] else "Off"
            cells[2].text = ", ".join(w["properties_written"][:4]) or "—"
            cells[3].text = str(w["action_count"])

    # -- workflow reference ------------------------------------------------
    doc.add_page_break()
    doc.add_heading("Workflow reference", level=1)
    doc.add_paragraph(
        "Every workflow with its enrollment logic, the properties it reads and writes, and a "
        "link into the HubSpot editor."
    )
    for w in sorted(workflows, key=lambda w: (not w["enabled"], w["name"].lower())):
        doc.add_heading(w["name"], level=2)
        p = doc.add_paragraph()
        _hyperlink(p, w["link"], "Open in HubSpot")
        p.add_run(f"   ·   ID {w['id']}   ·   ")
        p.add_run("Active" if w["enabled"] else "Turned off").bold = True
        p.add_run(f"   ·   Updated {_date(w['updated_at'])}")
        doc.add_paragraph(_strip_ticks(w["description"]))
        if w["properties_written"]:
            _kv(doc, "Writes", ", ".join(w["properties_written"]))
        if w["properties_read"]:
            _kv(doc, "Reads", ", ".join(w["properties_read"]))
        if w["workflows_triggered"]:
            names = [by_id[i]["name"] if i in by_id else i for i in w["workflows_triggered"]]
            _kv(doc, "Triggers", ", ".join(names))
        if w["lists_referenced"]:
            _kv(doc, "Lists referenced", ", ".join(w["lists_referenced"]))

    # -- property reference ------------------------------------------------
    doc.add_page_break()
    doc.add_heading("Property reference", level=1)
    doc.add_paragraph(
        "Sorted by how much of the portal depends on each property. Usage counts workflow "
        "writes and reads, form collection, and list segmentation — a property nothing "
        "references is a property nothing would miss."
    )
    t = doc.add_table(rows=1, cols=7)
    t.style = "Light Grid Accent 1"
    for i, head in enumerate(["Property", "Label", "Object", "Type", "Group", "Usage", "Referenced by"]):
        t.rows[0].cells[i].text = head
    for p in props:
        cells = t.add_row().cells
        _hyperlink(cells[0].paragraphs[0], p["link"], p["name"])
        cells[1].text = p.get("label") or "—"
        cells[2].text = p["object_type"]
        cells[3].text = p.get("field_type") or p.get("type") or "—"
        cells[4].text = p.get("group") or "—"
        cells[5].text = str(p["usage_score"])
        refs = []
        n_wf = len(p["written_by_workflows"]) + len(p["read_by_workflows"])
        if n_wf:
            refs.append(f"{n_wf} workflow(s)")
        if p["collected_by_forms"]:
            refs.append(f"{len(p['collected_by_forms'])} form(s)")
        if p["segmented_by_lists"]:
            refs.append(f"{len(p['segmented_by_lists'])} list(s)")
        cells[6].text = ", ".join(refs) or ("unused" if p["unused"] else "—")

    # -- remaining asset families -----------------------------------------
    doc.add_page_break()
    _asset_section(doc, "Lists", analysis["lists"],
                   ["List", "Type", "Size", "Segments on"],
                   lambda l: [l["name"], l.get("processing_type") or "—",
                              str(l.get("size") or "—"), ", ".join(l["properties"][:5]) or "—"],
                   lambda l: l["link"])
    _asset_section(doc, "Forms", analysis["forms"],
                   ["Form", "Type", "Fields", "Writes to"],
                   lambda f: [f["name"], f.get("type") or "—", str(f["field_count"]),
                              ", ".join(f["fields"][:6]) or "—"],
                   lambda f: f["link"])
    _asset_section(doc, "Marketing emails", analysis["marketing_emails"],
                   ["Email", "Subject", "State", "Updated"],
                   lambda e: [e["name"], e.get("subject") or "—", e.get("state") or "—",
                              _date(e.get("updated_at"))],
                   lambda e: e["link"])
    _asset_section(doc, "Landing & site pages",
                   analysis["landing_pages"] + analysis["site_pages"],
                   ["Page", "Kind", "URL", "State"],
                   lambda p: [p["name"], p["kind"], p.get("url") or "—", p.get("state") or "—"],
                   lambda p: p["link"])

    # -- caveats -----------------------------------------------------------
    doc.add_heading("Coverage & caveats", level=1)
    skipped = analysis.get("skipped") or []
    if skipped:
        doc.add_paragraph("These endpoints could not be read with the token supplied:")
        for s in skipped:
            doc.add_paragraph(f"{s['endpoint']} — {s.get('status') or ''} {str(s.get('reason'))[:160]}",
                              style="List Bullet")
    else:
        doc.add_paragraph("Every endpoint this audit queries returned data.")
    unknown = analysis.get("unknown_action_types") or {}
    if unknown:
        doc.add_paragraph(
            "Workflow action type IDs not in this tool's label table, reported by raw ID "
            "rather than guessed at: "
            + ", ".join(f"{k} (×{v})" for k, v in unknown.items())
        )
    doc.add_paragraph(
        "Workflow definitions come from HubSpot's v4 Automation API, which returns logic but "
        "not enrollment history — how many records a workflow has touched lives only in the "
        "HubSpot UI. Property references are resolved by scanning workflow, list, and form "
        "definitions for names matching this portal's own property schema, so a reference "
        "made through an integration or a custom-coded action may not appear."
    )

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path


def _asset_section(doc, title, items, headers, row_fn, link_fn) -> None:
    doc.add_heading(title, level=1)
    if not items:
        doc.add_paragraph("None found, or the endpoint was not reachable with this token.")
        return
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    for i, head in enumerate(headers):
        t.rows[0].cells[i].text = head
    for item in items:
        values = row_fn(item)
        cells = t.add_row().cells
        _hyperlink(cells[0].paragraphs[0], link_fn(item), values[0])
        for i, value in enumerate(values[1:], start=1):
            cells[i].text = str(value)


def _kv(doc, label: str, value: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(f"{label}: ").bold = True
    p.add_run(value)


def _hyperlink(paragraph, url: str, text: str) -> None:
    """python-docx has no hyperlink API; build the w:hyperlink run directly."""
    from docx.oxml.shared import OxmlElement, qn

    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0F5F63")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(color)
    props.append(underline)
    run.append(props)
    text_el = OxmlElement("w:t")
    text_el.text = text
    run.append(text_el)
    link.append(run)
    paragraph._p.append(link)


def _docx_styles(doc, Pt, RGBColor) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    for level, size, color in (("Title", 26, "12171A"), ("Heading 1", 17, "0F5F63"),
                               ("Heading 2", 13, "12171A")):
        try:
            style = styles[level]
        except KeyError:
            continue
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = level != "Title"
