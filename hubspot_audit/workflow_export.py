"""Read HubSpot's own workflow listing export and merge it into the snapshot.

The Automation API says what a workflow is *made of*; it says nothing about what the
workflow has been *doing*. The operational half — when it last acted on a record, how
many records went through it in the last seven days, how many sit in it right now, how
many actions are currently erroring — lives on the workflow CRM object, and the only
way to get that object out of HubSpot is a listing export.

So the audit accepts one as a file. Merging is by flow id, with an unambiguous-name
fallback, and the merge is always additive: the export can add activity to a workflow
the API already described, and it can add a workflow the API run never saw, but it
never overwrites a definition.

Column meanings, confirmed against the legacy Automation API on this portal:

    Enrolled unique   lifetime distinct records — equals contactCounts.enrolled
    Enrolled total    lifetime enrolments including re-enrolment, so >= unique
    Currently Enrolled  records inside the workflow right now
    Last action on    the last time an action executed — set for 513 workflows that
                      have never had a single issue, so it is execution recency and
                      not a rebadged error timestamp
"""

from __future__ import annotations

import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

# Export header -> the key we carry on the flow record. Anything not listed is dropped:
# folder ids, empty team columns and the always-false credit flag tell a reader nothing.
COLUMNS = {
    "Flow ID": "flow_id",
    "Name": "name",
    "On or Off": "enabled",
    "Object type": "object_type",
    "Description": "description",
    "Last action on": "last_action_at",
    "Enrolled last 7-days": "enrolled_7d",
    "Enrolled unique": "enrolled_unique",
    "Enrolled total": "enrolled_runs",
    "Currently Enrolled": "currently_enrolled",
    "Current Issue Count": "open_issues",
    "Last issue occurred on": "last_issue_at",
    "Action type": "action_types",
    "Trigger Type": "trigger_type",
    "Is re-enrollment enabled": "re_enrollment",
    "Used by other tools": "used_elsewhere",
    "Created in": "built_in",
    "Created on": "created_at",
    "Created by": "created_by",
    "Updated on": "updated_at",
    "Updated by": "updated_by",
    "Brands": "brand",
}

BOOL = {"enabled", "re_enrollment", "used_elsewhere"}
# Excel hands these back as floats, so the id has to be normalised the same way the
# counts are — "1870549122.0" matches no flow id in the world.
NUMBER = {"flow_id", "enrolled_7d", "enrolled_unique", "enrolled_runs",
          "currently_enrolled", "open_issues"}
DATE = {"last_action_at", "last_issue_at", "created_at", "updated_at"}

# HubSpot's own trigger labels, said the way a reader would say them.
TRIGGER = {
    "Filter criteria": "When a record matches the starting filters",
    "Events": "When a specific event happens",
    "Schedule": "On a schedule",
}

BUILT_IN = {
    "WORKFLOWS_APP": "Workflows tool",
    "WORKFLOWS_CLASSIC": "Workflows tool (classic editor)",
    "FORMS_APP": "Created from a form",
    "DEAL_PIPELINE_SETTINGS": "Created from deal pipeline settings",
}

_DATE_IN_NAME = re.compile(r"(\d{2})(\d{2})(\d{4})")


def _clean(value: Any, kind: str | None) -> Any:
    if value is None or value == "":
        return None
    if kind in BOOL:
        return str(value).strip().lower() == "true"
    if kind in NUMBER:
        try:
            return int(float(value))
        except (TypeError, ValueError):
            return None
    if kind in DATE:
        if isinstance(value, datetime):
            stamp = value if value.tzinfo else value.replace(tzinfo=timezone.utc)
            return stamp.isoformat()
        return str(value)
    text = str(value).strip()
    return text or None


def exported_at(path: Path, rows: list[dict[str, Any]] | None = None) -> str | None:
    """When the export was taken.

    HubSpot puts the date in the filename (…04092026…, DDMMYYYY), but the file gets
    renamed on its way into a repo. So fall back to the newest timestamp anywhere in
    the data, which is a true lower bound: the export cannot predate its own contents.
    """
    match = _DATE_IN_NAME.search(path.stem)
    if match:
        day, month, year = match.groups()
        try:
            return datetime(int(year), int(month), int(day), tzinfo=timezone.utc).date().isoformat()
        except ValueError:
            pass
    stamps = [str(row[key]) for row in (rows or []) for key in DATE if row.get(key)]
    return max(stamps)[:10] if stamps else None


def load(path: Path) -> list[dict[str, Any]]:
    """Every row of the export, as flat dicts using our own key names."""
    from openpyxl import load_workbook

    book = load_workbook(path, read_only=True, data_only=True)
    sheet = book.worksheets[0]
    # HubSpot writes no sheet dimension, so read-only mode reports a 1x1 sheet.
    sheet.reset_dimensions()

    rows = sheet.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(rows)]
    keys = [COLUMNS.get(h) for h in header]

    out: list[dict[str, Any]] = []
    for raw in rows:
        record: dict[str, Any] = {}
        for key, value in zip(keys, raw):
            if key:
                record[key] = _clean(value, key)
        if record.get("flow_id"):
            record["flow_id"] = str(record["flow_id"])
            record["action_types"] = _split_actions(record.get("action_types"))
            record["built_in"] = BUILT_IN.get(record.get("built_in") or "", record.get("built_in"))
            trigger = (record.get("trigger_type") or "").strip().strip("-").strip()
            record["trigger_type"] = TRIGGER.get(trigger, trigger) or None
            out.append(record)
    book.close()
    return out


def _split_actions(value: Any) -> list[str]:
    if not value:
        return []
    return [part.strip() for part in str(value).split(";") if part.strip()]


def merge(snapshot: dict[str, Any], path: Path) -> dict[str, Any]:
    """Attach export activity to the snapshot's workflows and append the ones it lacks.

    Returns a small report the deliverables quote, so no document has to claim a
    coverage number that was not actually measured.
    """
    rows = load(path)
    flows = snapshot.get("workflows") or []

    by_id = {str(f.get("id")): f for f in flows}
    by_name: dict[str, list[dict[str, Any]]] = {}
    for flow in flows:
        by_name.setdefault((flow.get("name") or "").strip().lower(), []).append(flow)

    matched: set[str] = set()
    added: list[dict[str, Any]] = []

    for row in rows:
        flow = by_id.get(row["flow_id"])
        if flow is None:
            # Fall back to the name, but only when it picks out exactly one workflow.
            candidates = by_name.get((row.get("name") or "").strip().lower(), [])
            flow = candidates[0] if len(candidates) == 1 else None
        if flow is None:
            added.append(_synthesize(row))
            continue
        _apply(flow, row)
        matched.add(str(flow.get("id")))

    for flow in flows:
        if str(flow.get("id")) not in matched:
            flow["_export_present"] = False

    flows.extend(added)
    snapshot["workflows"] = flows
    snapshot["workflow_export"] = {
        "exported_at": exported_at(path, rows),
        "rows": len(rows),
        "matched": len(matched),
        "added": len(added),
        "missing_from_export": len(flows) - len(added) - len(matched),
    }
    return snapshot["workflow_export"]


def _apply(flow: dict[str, Any], row: dict[str, Any]) -> None:
    """Copy activity onto a flow the API already described. Definitions are untouched."""
    flow["_export_present"] = True
    for key in ("last_action_at", "enrolled_7d", "currently_enrolled", "open_issues",
                "last_issue_at", "trigger_type", "built_in", "created_by", "updated_by",
                "brand", "used_elsewhere", "enrolled_runs"):
        if row.get(key) is not None:
            flow[f"_export_{key}"] = row[key]
    if row.get("action_types"):
        flow["_export_action_types"] = row["action_types"]

    # The export is the fresher read of the same two numbers the legacy API gave us.
    if row.get("enrolled_unique") is not None:
        flow["_enrolled_total"] = row["enrolled_unique"]
    if row.get("currently_enrolled") is not None:
        flow["_enrolled_active"] = row["currently_enrolled"]
    elif row.get("enrolled_unique") == 0:
        flow["_enrolled_active"] = 0

    # A workflow with no enrolment row at all has never run. Say zero, not unknown.
    if flow.get("_enrolled_total") is None and row.get("last_action_at") is None:
        flow["_enrolled_total"] = 0
        flow["_enrolled_active"] = 0


def _synthesize(row: dict[str, Any]) -> dict[str, Any]:
    """A workflow that exists in HubSpot but was created after the API snapshot.

    It gets a real record so it appears everywhere the others do, flagged as having no
    readable definition — which is exactly what `_describe_steps` already handles.
    """
    flow = {
        "id": row["flow_id"],
        "name": row.get("name") or "(unnamed)",
        "isEnabled": bool(row.get("enabled")),
        "objectTypeId": _OBJECT_TYPE_IDS.get(row.get("object_type") or "", ""),
        "flowType": "WORKFLOW",
        "description": row.get("description") or "",
        "createdAt": row.get("created_at"),
        "updatedAt": row.get("updated_at"),
        "_definition_available": False,
        "_export_only": True,
    }
    _apply(flow, row)
    return flow


_OBJECT_TYPE_IDS = {
    "CONTACT": "0-1",
    "COMPANY": "0-2",
    "DEAL": "0-3",
    "TICKET": "0-5",
    "COMMERCE_PAYMENT": "0-101",
    "SUBSCRIPTION": "0-69",
    "USER": "0-115",
}
