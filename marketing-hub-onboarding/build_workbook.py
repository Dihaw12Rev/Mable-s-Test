"""Build Marketing_Hub_Onboarding_Process.xlsx from process_data.py.

Run:  python3 build_workbook.py
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

import process_data as d

OUT = Path(__file__).with_name("Marketing_Hub_Onboarding_Process.xlsx")

FONT = "Arial"
NAVY = "1F3A5F"
TEAL = "2E7D7A"
PALE = "EEF3F7"
YELLOW = "FFF4C2"
BAR = "2E7D7A"

thin = Side(style="thin", color="D5DCE3")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

PHASE_NAME = {p["id"]: p["name"] for p in d.PHASES}


def style_header(ws, row, ncols, fill=NAVY, first_col=1):
    for c in range(first_col, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.alignment = CENTER
        cell.border = BORDER
    ws.row_dimensions[row].height = 30


def style_body(ws, first_row, last_row, ncols, first_col=1):
    for r in range(first_row, last_row + 1):
        for c in range(first_col, ncols + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = Font(name=FONT, size=10)
            cell.alignment = WRAP
            cell.border = BORDER


def title(ws, text, subtitle):
    ws["A1"] = text
    ws["A1"].font = Font(name=FONT, bold=True, size=14, color=NAVY)
    ws["A2"] = subtitle
    ws["A2"].font = Font(name=FONT, italic=True, size=10, color="555555")


def heading(ws, row, text, col=1):
    ws.cell(row=row, column=col, value=text)
    ws.cell(row=row, column=col).font = Font(name=FONT, bold=True, size=12, color=NAVY)


def widths(ws, values):
    for i, w in enumerate(values, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def table(ws, header_row, headers, rows, col_widths=None, freeze=True, fill=NAVY):
    for c, h in enumerate(headers, start=1):
        ws.cell(row=header_row, column=c, value=h)
    style_header(ws, header_row, len(headers), fill=fill)
    for r, row in enumerate(rows, start=header_row + 1):
        for c, v in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=v)
    last = header_row + len(rows)
    style_body(ws, header_row + 1, last, len(headers))
    if col_widths:
        widths(ws, col_widths)
    if freeze:
        ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(headers))}{last}"
    return last


def yellow(ws, cells):
    for ref in cells:
        ws[ref].fill = PatternFill("solid", fgColor=YELLOW)


def start_here(wb):
    ws = wb.active
    ws.title = "Start Here"
    title(ws, "RevGravy · Marketing Hub Onboarding Process (v2)",
          f"Built from the FigJam session and the answers from review. Day 0 = deal reaches 'Start Payment Received'. "
          f"Programs run {d.PROGRAM_MIN_DAYS}–{d.PROGRAM_DAYS} days; the plan is laid out for {d.PROGRAM_DAYS}.")
    rows = [
        ("Decisions", "Your answers, what each one changed, and the scope limits."),
        ("Roadmap", "Phases on a week-by-week timeline, mapped to the FigJam stickies."),
        ("Journey", "Start to end, lane by lane (customer, Phil, Architect, Ops, HubSpot Work), with when to mark tasks Done."),
        ("Task Template", "Every task by section. Set Program length (60–90) at the top and the due days compress."),
        ("Progress Measures", "The standard progress measures (Factories) every Marketing Hub program builds."),
        ("Requests", "How $250 requests work and the urgency classes."),
        ("After the Program", "6 months of free KPI monitoring and the paid offers."),
        ("Intake Form", "What we ask the customer after the sale."),
        ("Tech Setup", "Marketing Hub Pro & Enterprise setup checklist; each item belongs to a task."),
        ("Training", "Training modules by audience and timing."),
        ("Comms", "Every customer touchpoint, including the weekly update outline."),
        ("HubSpot Work Setup", "What to build in HubSpot Work: new-customer steps, task columns, board, views, requests, Customers table, agents."),
        ("Tool Fit", "The Work vs Projects vs Tickets comparison (decision: HubSpot Work only)."),
    ]
    last = table(ws, 4, ["Tab", "What's in it"], rows, [24, 110], freeze=False)
    r = last + 2
    last = table_at(ws, r, ["Role", "Responsibility"], list(d.ROLES.items()))
    r = last + 2
    heading(ws, r, "Legend")
    ws.cell(row=r + 1, column=1, value="Yellow cell").fill = PatternFill("solid", fgColor=YELLOW)
    ws.cell(row=r + 1, column=2, value="Edit me: inputs and answers.")
    ws.cell(row=r + 2, column=1, value="■ bars")
    ws.cell(row=r + 2, column=2, value="Formulas on the Start/End day columns. Change the days and the bars move.")
    for rr in (r + 1, r + 2):
        for c in (1, 2):
            ws.cell(row=rr, column=c).font = Font(name=FONT, size=10)


def table_at(ws, row, headers, rows, fill=NAVY):
    """A table that starts lower on a sheet (no freeze, no filter)."""
    for c, h in enumerate(headers, start=1):
        ws.cell(row=row, column=c, value=h)
    style_header(ws, row, len(headers), fill=fill)
    for r, vals in enumerate(rows, start=row + 1):
        for c, v in enumerate(vals, start=1):
            ws.cell(row=r, column=c, value=v)
    last = row + len(rows)
    style_body(ws, row + 1, last, len(headers))
    return last


def decisions(wb):
    ws = wb.create_sheet("Decisions")
    title(ws, "Decisions", "Answers from the review and what each one changed. Open questions at the bottom: answer in yellow.")
    last = table(ws, 4, ["#", "Topic", "Answer", "What changed"],
                 [(i, t, a, w) for i, (t, a, w) in enumerate(d.DECISIONS, start=1)], [5, 26, 40, 80], freeze=False)
    r = last + 2
    heading(ws, r, "Scope limits for migration and website work (suggested)")
    last = table_at(ws, r + 1, ["#", "Area", "Limit", ""], [(i, a, l, "") for i, (a, l) in enumerate(d.SCOPE_LIMITS, start=1)], fill=TEAL)
    r = last + 2
    heading(ws, r, "Still open")
    if not d.OPEN_QUESTIONS:
        ws.cell(row=r + 1, column=1, value="Nothing open. Every question has been answered.").font = Font(name=FONT, size=10)
        return
    last = table_at(ws, r + 1, ["#", "Topic", "Question", "What it changes", "Answer"],
                    [(i, t, q, w, "") for i, (t, q, w) in enumerate(d.OPEN_QUESTIONS, start=1)])
    ws.column_dimensions["E"].width = 45
    for rr in range(r + 2, last + 1):
        ws.cell(row=rr, column=5).fill = PatternFill("solid", fgColor=YELLOW)


def gantt_columns(ws, header_row, first_col, start_col, end_col, first_data, last_data):
    """Week buckets from pre-sale to after the program; formulas mark where each row is active."""
    buckets = [("Pre", -60, -1)] + [(f"W{w}", w * 7, w * 7 + 6) for w in range(13)] + [("After", 91, 400)]
    for i, (label, lo, hi) in enumerate(buckets):
        col = first_col + i
        letter = get_column_letter(col)
        ws.cell(row=header_row - 2, column=col, value=lo)
        ws.cell(row=header_row - 1, column=col, value=hi)
        ws.cell(row=header_row, column=col, value=label)
        for rr in (header_row - 2, header_row - 1):
            ws.cell(row=rr, column=col).font = Font(name=FONT, size=8, color="999999")
            ws.cell(row=rr, column=col).alignment = CENTER
        ws.column_dimensions[letter].width = 5.5
        for r in range(first_data, last_data + 1):
            cell = ws.cell(row=r, column=col,
                           value=f'=IF(AND(${start_col}{r}<={letter}${header_row - 1},'
                                 f'${end_col}{r}>={letter}${header_row - 2}),"■","")')
            cell.alignment = CENTER
            cell.border = BORDER
            cell.font = Font(name=FONT, size=10, color=BAR)
    last_letter = get_column_letter(first_col + len(buckets) - 1)
    ws.conditional_formatting.add(
        f"{get_column_letter(first_col)}{first_data}:{last_letter}{last_data}",
        CellIsRule(operator="equal", formula=['"■"'], fill=PatternFill("solid", fgColor=BAR)))
    return first_col + len(buckets) - 1


def roadmap(wb):
    ws = wb.create_sheet("Roadmap")
    title(ws, "Roadmap", f"{d.PROGRAM_DAYS}-day plan. Grey numbers above the weeks are each week's day range.")
    headers = ["Phase", "Name", "FigJam stage(s)", "Start day", "End day", "Goal", "What the client sees", "Done when"]
    hr = 6
    rows = [(p["id"], p["name"], p["figjam"], p["start"], p["end"], p["goal"], p["client_sees"], p["exit"]) for p in d.PHASES]
    last = table(ws, hr, headers, rows, [7, 20, 30, 8, 8, 40, 40, 40])
    for r in range(hr + 1, last + 1):
        for c in (4, 5):
            ws.cell(row=r, column=c).fill = PatternFill("solid", fgColor=YELLOW)
            ws.cell(row=r, column=c).alignment = CENTER
    lastcol = gantt_columns(ws, hr, len(headers) + 1, "D", "E", hr + 1, last)
    style_header(ws, hr, lastcol)
    ws.auto_filter.ref = None
    ws.freeze_panes = "C7"


def task_template(wb):
    ws = wb.create_sheet("Task Template")
    title(ws, "Task Template", "Grouped by section (phase). Set Program length in yellow; 'Due day (this program)' compresses.")
    ws["A3"], ws["C3"] = "Program length (days)", d.PROGRAM_DAYS
    ws["D3"], ws["E3"] = "Plan laid out for", d.PROGRAM_DAYS
    ws["F3"], ws["G3"] = "Fixed until day", d.FIXED_UNTIL_DAY
    for ref in ("A3", "D3", "F3"):
        ws[ref].font = Font(name=FONT, bold=True, size=10)
    for ref in ("C3", "E3", "G3"):
        ws[ref].font = Font(name=FONT, size=10)
        ws[ref].alignment = CENTER
    yellow(ws, ["C3"])
    dv_len = DataValidation(type="whole", operator="between", formula1=str(d.PROGRAM_MIN_DAYS), formula2=str(d.PROGRAM_DAYS))
    ws.add_data_validation(dv_len)
    dv_len.add("C3")

    headers = ["ID", "Section", "Task", "Definition of done", "Owner", "Client-facing", "Start day", "Due day (90-day plan)",
               "Due day (this program)", "Repeats", "Depends on", "Lives in"]
    hr = 6
    rows = [(tid, f"{ph} · {PHASE_NAME[ph]}", task, det, owner, cf, s, e, None, rep, dep, where)
            for tid, ph, task, det, owner, cf, s, e, dep, where, rep in d.TASKS]
    last = table(ws, hr, headers, rows, [6, 20, 32, 55, 11, 9, 7, 9, 9, 13, 11, 26])
    for r in range(hr + 1, last + 1):
        ws.cell(row=r, column=9, value=(
            f'=IF(OR(LEFT($B{r},2)="P0",$H{r}<=$G$3),$H{r},'
            f'IF(LEFT($B{r},2)="P9",$H{r}-$E$3+$C$3,ROUND($H{r}*$C$3/$E$3,0)))'))
        for c in (6, 7, 8, 9):
            ws.cell(row=r, column=c).alignment = CENTER
        for c in (7, 8):
            ws.cell(row=r, column=c).fill = PatternFill("solid", fgColor=YELLOW)
        ws.cell(row=r, column=9).fill = PatternFill("solid", fgColor=PALE)
    ws.conditional_formatting.add(
        f"F{hr + 1}:F{last}", CellIsRule(operator="equal", formula=['"Yes"'], fill=PatternFill("solid", fgColor="DCEFEA")))
    dv = DataValidation(type="list", formula1='"' + ",".join(d.ROLES) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"E{hr + 1}:E{last}")
    lastcol = gantt_columns(ws, hr, len(headers) + 1, "G", "H", hr + 1, last)
    style_header(ws, hr, lastcol)
    ws.auto_filter.ref = f"A{hr}:{get_column_letter(len(headers))}{last}"
    ws.freeze_panes = "D7"

    r = last + 3
    heading(ws, r, "Tasks by owner", col=3)
    for c, h in enumerate(["Owner", "# tasks", "Client-facing"], start=3):
        ws.cell(row=r + 1, column=c, value=h)
    style_header(ws, r + 1, 5, first_col=3)
    for i, role in enumerate(d.ROLES, start=r + 2):
        ws.cell(row=i, column=3, value=role)
        ws.cell(row=i, column=4, value=f"=COUNTIF($E${hr + 1}:$E${last},C{i})")
        ws.cell(row=i, column=5, value=f'=COUNTIFS($E${hr + 1}:$E${last},C{i},$F${hr + 1}:$F${last},"Yes")')
    tot = r + 2 + len(d.ROLES)
    ws.cell(row=tot, column=3, value="Total")
    ws.cell(row=tot, column=4, value=f"=SUM(D{r + 2}:D{tot - 1})")
    ws.cell(row=tot, column=5, value=f"=SUM(E{r + 2}:E{tot - 1})")
    style_body(ws, r + 2, tot, 5, first_col=3)
    for c in (3, 4, 5):
        ws.cell(row=tot, column=c).font = Font(name=FONT, size=10, bold=True)


def simple_tab(wb, name, subtitle, headers, rows, col_widths):
    ws = wb.create_sheet(name)
    title(ws, name, subtitle)
    table(ws, 4, headers, rows, col_widths)
    return ws


def requests(wb):
    ws = wb.create_sheet("Requests")
    title(ws, "Requests", "$250 per request. Recommended rules and urgency classes.")
    heading(ws, 4, "Rules")
    for i, rule in enumerate(d.REQUEST_RULES, start=5):
        ws.cell(row=i, column=1, value=f"{i - 4}. {rule}").font = Font(name=FONT, size=10)
    r = 6 + len(d.REQUEST_RULES)
    table_at(ws, r, ["Urgency", "What it is", "Response and turnaround", "Price"], d.REQUEST_CLASSES)
    widths(ws, [14, 60, 48, 40])


def work_setup(wb):
    ws = wb.create_sheet("HubSpot Work Setup")
    title(ws, "HubSpot Work Setup", "What the HubSpot Work prompt builds. Decision: onboarding lives only in HubSpot Work.")
    widths(ws, [26, 26, 70])
    r = 4
    heading(ws, r, "Where each part lives")
    r = table_at(ws, r + 1, ["Tool", "", "Job"], [(t, "", j) for t, j in d.WHERE_IT_LIVES], fill=TEAL) + 2
    heading(ws, r, "Setting up a new customer (all customers share the same tables)")
    r = table_at(ws, r + 1, ["Step", "", "What to do"], [(i, "", st) for i, st in enumerate(d.NEW_CUSTOMER_STEPS, start=1)], fill=TEAL) + 2
    heading(ws, r, "Task columns (template Tasks table and the shared Customer Tasks table, which adds a Customer link)")
    r = table_at(ws, r + 1, ["Column", "Type", "Notes"], d.WORK_TASK_COLUMNS) + 2
    heading(ws, r, "Board statuses")
    r = table_at(ws, r + 1, ["Status", "", ""], [(s, "", "") for s in d.WORK_STATUSES]) + 2
    heading(ws, r, "Views")
    r = table_at(ws, r + 1, ["View", "", "Setup"], [(v, "", s) for v, s in d.WORK_VIEWS]) + 2
    heading(ws, r, "Requests table (shared, each row linked to its customer)")
    r = table_at(ws, r + 1, ["Column", "Type", ""], [(c, t, "") for c, t in d.WORK_REQUEST_COLUMNS]) + 2
    heading(ws, r, "Updates table (drafts the agents write for the Architect)")
    r = table_at(ws, r + 1, ["Column", "Type", ""], [(c, t, "") for c, t in d.WORK_UPDATE_COLUMNS]) + 2
    heading(ws, r, "Customers table, one row per customer")
    r = table_at(ws, r + 1, ["Column", "Type", ""], [(c, t, "") for c, t in d.WORK_PORTFOLIO_COLUMNS]) + 2
    r = table_at(ws, r, ["Portfolio view", "", "Setup"], [(v, "", s) for v, s in d.WORK_PORTFOLIO_VIEWS]) + 2
    heading(ws, r, "Agents (draft, fill fields or flag only; none sends anything to a customer). Full prompts: hubspot-work-agent-prompts.md")
    table_at(ws, r + 1, ["Agent", "Trigger", "What it does"], [(a["name"], a["trigger"], a["summary"]) for a in d.WORK_AGENTS], fill=TEAL)


def journey(wb):
    ws = wb.create_sheet("Journey")
    ex = d.EXAMPLE
    title(ws, "Journey", f"Worked example: {ex['customer']} ({ex['note'].lower()}), {ex['tier']}, {ex['program_days']} days, Day 0 = {ex['day0']}.")
    names = {t[0]: t[2] for t in d.TASKS}
    lane_keys = [k for k, n, h in d.LANES]
    headers = ["Stage", "Day", "Moment"] + [n for k, n, h in d.LANES] + ["Mark Done in Work", "Note"]
    rows = []
    for m in d.JOURNEY:
        days = "Before Day 0" if m["end"] < 0 else (f"Day {m['start']}" if m["start"] == m["end"] else f"Day {m['start']}–{m['end']}")
        done = ", ".join(m["done"]) + (("; Won't do: " + ", ".join(m["wont"])) if m.get("wont") else "")
        rows.append([m["stage"], days, m["title"]] + ["\n".join("• " + x for x in m[k]) or "—" for k in lane_keys] + [done or "—", m.get("note", "")])
    last = table(ws, 4, headers, rows, [14, 12, 20, 40, 34, 44, 30, 40, 20, 30], freeze=False)
    ws.freeze_panes = "D5"
    r = last + 2
    heading(ws, r, "When to change a task's status")
    r = table_at(ws, r + 1, ["Status", "When"], d.STATUS_RULES) + 2
    heading(ws, r, "How a request moves")
    r = table_at(ws, r + 1, ["Step", "What happens"], d.REQUEST_FLOW) + 2
    heading(ws, r, "A normal week")
    r = table_at(ws, r + 1, ["When", "Who", "What happens"], d.WEEKLY_RHYTHM) + 2
    heading(ws, r, "When things don't go to plan")
    table_at(ws, r + 1, ["Situation", "What to do"], d.EXCEPTIONS)


def tool_fit(wb):
    ws = wb.create_sheet("Tool Fit")
    title(ws, "Tool Fit: HubSpot Work vs Projects vs Tickets",
          "Decision: onboarding lives only in HubSpot Work. Scores kept as the record (0 = poor, 1 = workable, 2 = strong).")
    headers = ["Criterion", "Work", "HubSpot Work", "Projects", "Projects object", "Tickets", "Tickets"]
    rows = [(c, w[0], w[1], p[0], p[1], t[0], t[1]) for c, w, p, t in d.TOOL_FIT]
    last = table(ws, 4, headers, rows, [30, 7, 38, 9, 34, 8, 32], freeze=False)
    tr = last + 1
    ws.cell(row=tr, column=1, value=f"Total (max {2 * len(d.TOOL_FIT)})")
    for c in (2, 4, 6):
        L = get_column_letter(c)
        ws.cell(row=tr, column=c, value=f"=SUM({L}5:{L}{last})")
    for c in range(1, 8):
        cell = ws.cell(row=tr, column=c)
        cell.font = Font(name=FONT, bold=True, size=10)
        cell.fill = PatternFill("solid", fgColor=PALE)
        cell.border = BORDER
        cell.alignment = CENTER if c in (2, 4, 6) else WRAP
    for r in range(5, last + 1):
        for c in (2, 4, 6):
            ws.cell(row=r, column=c).alignment = CENTER


def main():
    wb = Workbook()
    start_here(wb)
    decisions(wb)
    roadmap(wb)
    journey(wb)
    task_template(wb)
    simple_tab(wb, "Progress Measures", "Name each one '<Client> 1.1 Visitor to Lead', like the existing Factory records.",
               ["ID", "Progress measure", "When", "What it measures", "HubSpot report", "Built by tasks"],
               d.FACTORIES, [6, 28, 30, 50, 36, 16])
    requests(wb)
    simple_tab(wb, "After the Program", "Starts the day the program ends. Prices in yellow are still to be decided.",
               ["Offer", "Price", "What's included"], d.POST_PROGRAM, [28, 44, 90])
    simple_tab(wb, "Intake Form", "Post-sale intake. 'Captured pre-sale' items are pre-filled and only confirmed.",
               ["Section", "Question", "Field type", "Required", "Captured pre-sale?", "Why we ask", "Saved in"],
               d.INTAKE, [16, 50, 34, 9, 14, 36, 24])
    simple_tab(wb, "Tech Setup", "Marketing Hub Professional and Enterprise. 'Ent only' items apply to Enterprise customers.",
               ["Area", "Item", "Details", "Tier", "Task ID", "Owner"], d.TECH_SETUP, [14, 36, 60, 10, 8, 22])
    simple_tab(wb, "Training", "Each module is short and timed to when the feature ships.",
               ["Module", "Audience", "Content", "Format", "When", "Length"], d.TRAINING, [30, 26, 50, 30, 30, 10])
    simple_tab(wb, "Comms", "Every customer-facing touchpoint, in order.",
               ["Touchpoint", "Timing", "Channel", "Owner", "Content"], d.COMMS, [26, 24, 22, 24, 80])
    work_setup(wb)
    tool_fit(wb)

    ws = wb["After the Program"]
    for r in range(5, 5 + len(d.POST_PROGRAM)):
        if ws.cell(row=r, column=2).value != "Free":
            ws.cell(row=r, column=2).fill = PatternFill("solid", fgColor=YELLOW)
    for ws in wb.worksheets:
        ws.sheet_properties.tabColor = NAVY if ws.title in ("Start Here", "Decisions") else TEAL
    wb.save(OUT)
    print(f"Saved {OUT}")


if __name__ == "__main__":
    main()
